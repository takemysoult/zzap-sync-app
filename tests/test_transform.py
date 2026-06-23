from openpyxl import load_workbook

from conftest import make_rows
from engine.transform import (apply_exclusions, build_xlsx, clean_rows,
                              parse_exclusions)


def test_clean_rows_drops_blank_number():
    cleaned = clean_rows(make_rows())
    assert len(cleaned) == 2
    assert all(r.number for r in cleaned)


def test_parse_exclusions_normalizes_and_skips_comments():
    text = "a-1\n# comment\n  B-2  \n\n"
    assert parse_exclusions(text) == {"A-1", "B-2"}


def test_apply_exclusions_is_case_insensitive():
    # Real pipeline: parse normalizes the list (upper), apply normalizes each row's
    # article before comparing — so a lowercase "a-1" entry still excludes row "A-1".
    excluded = parse_exclusions("a-1")
    kept = apply_exclusions(clean_rows(make_rows()), excluded)
    assert [r.number for r in kept] == ["B-2"]


def test_build_xlsx_no_header_default_layout(tmp_path):
    out = build_xlsx(clean_rows(make_rows()), tmp_path / "price.xlsx", include_header=False)
    ws = load_workbook(out).active
    data = list(ws.iter_rows(values_only=True))
    assert len(data) == 2          # no header row
    assert ws.max_column == 5
    # order: producer | number | name | quantity | price
    assert data[0][0] == "BrandA"
    assert data[0][1] == "A-1"
    assert data[0][2] == "Деталь 1"
    assert float(data[0][3]) == 5.0
    assert float(data[0][4]) == 100.0


def test_build_xlsx_with_header_writes_russian_titles(tmp_path):
    out = build_xlsx(clean_rows(make_rows()), tmp_path / "h.xlsx", include_header=True)
    ws = load_workbook(out).active
    header = list(ws.iter_rows(values_only=True))[0]
    assert header == ("Производитель", "Номер", "Наименование", "Количество", "Цена")


def test_build_xlsx_custom_column_order(tmp_path):
    cols = {"number": 1, "price": 2, "producer": 3, "name": 4, "quantity": 5}
    out = build_xlsx(clean_rows(make_rows()), tmp_path / "c.xlsx",
                     include_header=False, columns=cols)
    ws = load_workbook(out).active
    first = list(ws.iter_rows(values_only=True))[0]
    assert first[0] == "A-1"        # col1 = number
    assert float(first[1]) == 100.0  # col2 = price
    assert first[2] == "BrandA"     # col3 = producer
    assert first[3] == "Деталь 1"   # col4 = name
    assert float(first[4]) == 5.0   # col5 = quantity
